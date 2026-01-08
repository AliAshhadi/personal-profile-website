import io
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook

from app.utils import is_missing, normalize_phone, slugify_label


def expected_columns(rows: List[dict]) -> List[str]:
    base_cols = ["name", "subtitle1", "subtitle2", "profileImageFile"]
    dynamic_cols: List[str] = []
    for idx, row in enumerate(rows):
        slug = slugify_label(row.get("label", ""), f"row{idx + 1}")
        row_type = row.get("type", "link")
        if row_type in {"phone", "email"}:
            dynamic_cols.append(f"{slug}Label")
        else:
            dynamic_cols.append(f"{slug}Href")
    return base_cols + dynamic_cols


def build_template_workbook(rows: List[dict]) -> bytes:
    """
    Create an in-memory XLSX template with dynamic columns based on the row config.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    headers = expected_columns(rows)
    ws.append(headers)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_workbook(file_stream, rows: List[dict]) -> Tuple[List[dict], List[str]]:
    """
    Parse the first sheet of an XLSX file based on the provided row configuration.
    Returns (people, warnings).
    """
    wb = load_workbook(file_stream, data_only=True)
    sheet = wb.worksheets[0]
    warnings: List[str] = []
    if sheet.max_row < 2:
        return [], ["هیچ داده‌ای در فایل اکسل یافت نشد."]

    header_cells = [cell.value for cell in sheet[1]]
    headers = [_clean_text(h).lower() for h in header_cells]
    header_map = {name: idx for idx, name in enumerate(headers)}

    expected = expected_columns(rows)
    for column in expected:
        if column.lower() not in header_map:
            warnings.append(f"ستون مورد انتظار «{column}» یافت نشد.")

    people: List[dict] = []
    for row_idx in range(2, sheet.max_row + 1):
        row_values = sheet[row_idx]
        def get(col_name: str):
            idx = header_map.get(col_name.lower())
            if idx is None or idx >= len(row_values):
                return ""
            return _clean_text(row_values[idx].value)

        name = get("name")
        subtitle1 = get("subtitle1")
        subtitle2 = get("subtitle2")
        profile_image_file = get("profileImageFile")
        if is_missing(name) and is_missing(subtitle1) and is_missing(subtitle2):
            # skip empty row
            continue
        if is_missing(name):
            warnings.append(f"ردیف {row_idx}: ستون name خالی است. این ردیف نادیده گرفته شد.")
            continue

        person_rows: Dict[str, dict] = {}
        for idx, row_def in enumerate(rows):
            slug = slugify_label(row_def.get("label", ""), f"row{idx + 1}")
            row_type = row_def.get("type", "link")
            if row_type in {"phone", "email"}:
                label = get(f"{slug}Label")
                if is_missing(label):
                    continue
                value = normalize_phone(label) if row_type == "phone" else label
                person_rows[slug] = {
                    "value": label,
                    "normalized": value,
                }
            else:
                href = get(f"{slug}Href")
                if is_missing(href):
                    continue
                person_rows[slug] = {
                    "href": href,
                }

        people.append(
            {
                "name": name,
                "subtitle1": subtitle1,
                "subtitle2": subtitle2,
                "profileImageFile": profile_image_file,
                "rowValues": person_rows,
            }
        )

    return people, warnings
